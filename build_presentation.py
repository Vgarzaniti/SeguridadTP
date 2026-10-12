import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def create_deck():
    prs = Presentation()
    # 16:9 Widescreen dimensions
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    blank_layout = prs.slide_layouts[6]

    # Palette
    COLOR_BG = RGBColor(11, 17, 32)         # #0B1120 Dark Navy Slate
    COLOR_CARD = RGBColor(19, 29, 51)       # #131D33 Card Background
    COLOR_BORDER = RGBColor(36, 50, 77)     # #24324D Border
    COLOR_CYAN = RGBColor(6, 182, 212)      # #06B6D4 Accent Cyan
    COLOR_BLUE = RGBColor(59, 130, 246)     # #3B82F6 Blue
    COLOR_EMERALD = RGBColor(16, 185, 129)  # #10B981 Emerald
    COLOR_CRIMSON = RGBColor(239, 68, 68)   # #EF4444 Crimson
    COLOR_AMBER = RGBColor(245, 158, 11)    # #F59E0B Amber
    COLOR_PURPLE = RGBColor(139, 92, 246)   # #8B5CF6 Purple
    COLOR_WHITE = RGBColor(241, 245, 249)   # #F1F5F9 Heading
    COLOR_TEXT = RGBColor(203, 213, 225)    # #CBD5E1 Body text
    COLOR_MUTED = RGBColor(148, 163, 184)   # #94A3B8 Muted

    def apply_bg(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = COLOR_BG
        bg.line.fill.background() # no border
        return bg

    def add_header(slide, title, category_tag, speaker_tag, tag_color=COLOR_CYAN):
        # Speaker pill badge
        spk_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.4), Inches(3.2), Inches(0.4))
        spk_box.fill.solid()
        spk_box.fill.fore_color.rgb = COLOR_CARD
        spk_box.line.color.rgb = tag_color
        spk_box.line.width = Pt(1)
        tf_spk = spk_box.text_frame
        tf_spk.word_wrap = True
        p_spk = tf_spk.paragraphs[0]
        p_spk.text = speaker_tag.upper()
        p_spk.font.size = Pt(10.5)
        p_spk.font.bold = True
        p_spk.font.color.rgb = tag_color
        p_spk.alignment = PP_ALIGN.CENTER

        # Category tag
        if category_tag:
            cat_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(4.2), Inches(0.4), Inches(3.5), Inches(0.4))
            cat_box.fill.solid()
            cat_box.fill.fore_color.rgb = COLOR_CARD
            cat_box.line.color.rgb = COLOR_BORDER
            cat_box.line.width = Pt(1)
            tf_cat = cat_box.text_frame
            p_cat = tf_cat.paragraphs[0]
            p_cat.text = category_tag.upper()
            p_cat.font.size = Pt(10)
            p_cat.font.bold = True
            p_cat.font.color.rgb = COLOR_MUTED
            p_cat.alignment = PP_ALIGN.CENTER

        # Title
        t_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.85), Inches(11.7), Inches(0.8))
        tf_t = t_box.text_frame
        tf_t.word_wrap = True
        p_t = tf_t.paragraphs[0]
        p_t.text = title
        p_t.font.size = Pt(24)
        p_t.font.bold = True
        p_t.font.color.rgb = COLOR_WHITE

    def add_card(slide, left, top, width, height, title, items, badge=None, accent_color=COLOR_CYAN):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
        card.fill.solid()
        card.fill.fore_color.rgb = COLOR_CARD
        card.line.color.rgb = COLOR_BORDER
        card.line.width = Pt(1)

        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.25)
        tf.margin_right = Inches(0.25)
        tf.margin_top = Inches(0.22)
        tf.margin_bottom = Inches(0.22)

        # Title
        p0 = tf.paragraphs[0]
        p0.text = title
        p0.font.size = Pt(15)
        p0.font.bold = True
        p0.font.color.rgb = accent_color
        p0.space_after = Pt(8)

        if badge:
            p_badge = tf.add_paragraph()
            p_badge.text = f"[{badge}]"
            p_badge.font.size = Pt(9.5)
            p_badge.font.bold = True
            p_badge.font.color.rgb = COLOR_MUTED
            p_badge.space_after = Pt(8)

        for it in items:
            p = tf.add_paragraph()
            p.text = f"• {it}"
            p.font.size = Pt(12)
            p.font.color.rgb = COLOR_TEXT
            p.space_after = Pt(6)

        return card

    # =========================================================================
    # SLIDE 1: PORTADA
    # =========================================================================
    s1 = prs.slides.add_slide(blank_layout)
    apply_bg(s1)

    # Accent decorative bar
    bar = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.8), Inches(0.15), Inches(3.8))
    bar.fill.solid()
    bar.fill.fore_color.rgb = COLOR_CYAN
    bar.line.fill.background()

    # Title box
    t_box = s1.shapes.add_textbox(Inches(1.2), Inches(1.6), Inches(11.0), Inches(2.2))
    tf = t_box.text_frame
    tf.word_wrap = True
    p1 = tf.paragraphs[0]
    p1.text = "OWASP Top 10: 2025"
    p1.font.size = Pt(44)
    p1.font.bold = True
    p1.font.color.rgb = COLOR_WHITE

    p2 = tf.add_paragraph()
    p2.text = "Análisis de Riesgos Críticos en Aplicaciones Web & Mitigación Práctica"
    p2.font.size = Pt(20)
    p2.font.color.rgb = COLOR_CYAN
    p2.space_before = Pt(8)

    p3 = tf.add_paragraph()
    p3.text = "Basado en el estudio de 2.8M de aplicaciones y 248 debilidades de software (CWE)"
    p3.font.size = Pt(13)
    p3.font.color.rgb = COLOR_MUTED
    p3.space_before = Pt(8)

    # Team Info Box
    card_team = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.2), Inches(4.3), Inches(10.8), Inches(2.0))
    card_team.fill.solid()
    card_team.fill.fore_color.rgb = COLOR_CARD
    card_team.line.color.rgb = COLOR_BORDER
    card_team.line.width = Pt(1)
    tf_team = card_team.text_frame
    tf_team.margin_left = Inches(0.3)
    tf_team.margin_top = Inches(0.2)

    p_t1 = tf_team.paragraphs[0]
    p_t1.text = "PRESENTACIÓN GRUPAL (5 EXPOSITORES):"
    p_t1.font.size = Pt(12)
    p_t1.font.bold = True
    p_t1.font.color.rgb = COLOR_AMBER

    p_t2 = tf_team.add_paragraph()
    p_t2.text = "• Orador 1: Introducción, Metodología y Estado de la Seguridad\n• Orador 2: A01 (Control de Acceso) & A02 (Configuración Insegura)\n• Orador 3: A03 (Cadena de Suministro) & A04 (Fallas Criptográficas)\n• Orador 4: A05 (Inyección) & A06 (Diseño Inseguro)\n• Orador 5: A07 a A10 (Autenticación, Integridad, Logs, Excepciones) & Demo del Simulador"
    p_t2.font.size = Pt(11.5)
    p_t2.font.color.rgb = COLOR_TEXT
    p_t2.space_before = Pt(4)

    # =========================================================================
    # SLIDE 2: ESTRUCTURA & AGENDA
    # =========================================================================
    s2 = prs.slides.add_slide(blank_layout)
    apply_bg(s2)
    add_header(s2, "Agenda de la Presentación & Distribución de Roles", "Metodología de Exposición", "Equipo Completo", COLOR_BLUE)

    add_card(s2, 0.8, 1.8, 2.2, 5.0, "Orador 1", [
        "Qué es OWASP",
        "Espacio del problema (Web/APIs)",
        "Riesgo vs Vulnerabilidad",
        "Metodología edición 2025",
        "Comparativa 2021 vs 2025"
    ], "Bloque Introductorio", COLOR_BLUE)

    add_card(s2, 3.2, 1.8, 2.2, 5.0, "Orador 2", [
        "A01: Broken Access Control",
        "Absorción de SSRF",
        "100% incidencia en apps",
        "A02: Security Misconfiguration",
        "Ascenso al puesto #2",
        "Impacto estructural"
    ], "Top 1 y 2", COLOR_CRIMSON)

    add_card(s2, 5.6, 1.8, 2.2, 5.0, "Orador 3", [
        "A03: Software Supply Chain",
        "Nueva categoría 2025",
        "Caso SolarWinds & SBOM",
        "A04: Cryptographic Failures",
        "Algoritmos obsoletos",
        "Fugas de claves y TLS"
    ], "Top 3 y 4", COLOR_AMBER)

    add_card(s2, 8.0, 1.8, 2.2, 5.0, "Orador 4", [
        "A05: Injection (SQLi/XSS)",
        "+44k CVEs históricos",
        "Prepared Statements",
        "A06: Insecure Design",
        "Diseño vs Implementación",
        "Modelado de Amenazas"
    ], "Top 5 y 6", COLOR_PURPLE)

    add_card(s2, 10.4, 1.8, 2.2, 5.0, "Orador 5", [
        "A07: Authentication Failures",
        "A08: Integridad Software/Datos",
        "A09: Registro y Alertas",
        "A10: Condiciones Excepcionales",
        "Conclusiones Generales",
        "Demo Simulador Web"
    ], "Top 7-10 + Demo", COLOR_EMERALD)

    # =========================================================================
    # SLIDE 3: ORADOR 1 - INTRODUCCIÓN Y ESPACIO DEL PROBLEMA
    # =========================================================================
    s3 = prs.slides.add_slide(blank_layout)
    apply_bg(s3)
    add_header(s3, "Espacio del Problema: La Superficie de Ataque Moderna", "Fundamentos de Seguridad", "Orador 1")

    add_card(s3, 0.8, 1.8, 3.7, 5.0, "¿Qué es OWASP?", [
        "Open Worldwide Application Security Project.",
        "Organización global sin fines de lucro enfocada en la seguridad del software.",
        "El Top 10 es el estándar de facto adoptado por la industria internacional.",
        "Objetivo: Educar a desarrolladores y arquitectos en mitigación preventiva."
    ], "Estándar Global", COLOR_CYAN)

    add_card(s3, 4.8, 1.8, 3.7, 5.0, "El Espacio del Problema", [
        "Aplicaciones web acopladas a APIs y microservicios expuestos a Internet.",
        "Dependencia masiva de paquetes de terceros y repositorios públicos.",
        "Infraestructura cloud distribuida con pipelines CI/CD automatizados.",
        "Cada capa adicional expande exponencialmente la superficie de ataque."
    ], "Ecosistema Actual", COLOR_BLUE)

    add_card(s3, 8.8, 1.8, 3.7, 5.0, "Riesgo vs Vulnerabilidad", [
        "Vulnerabilidad: Debilidad o fallo técnico en el código o configuración.",
        "Riesgo = Vector de Ataque + Facilidad de Explotación + Impacto al Negocio.",
        "El OWASP Top 10 prioriza Riesgos: el daño real potencial a datos y operaciones.",
        "No todo bug técnico representa un riesgo crítico."
    ], "Diferenciación Clave", COLOR_EMERALD)

    # =========================================================================
    # SLIDE 4: ORADOR 1 - CONSTRUCCIÓN Y NOVEDADES 2025
    # =========================================================================
    s4 = prs.slides.add_slide(blank_layout)
    apply_bg(s4)
    add_header(s4, "Construcción del OWASP Top 10:2025 y Grandes Cambios", "Metodología y Datos", "Orador 1")

    add_card(s4, 0.8, 1.8, 5.7, 5.0, "Metodología Basada en Evidencia", [
        "Evaluación masiva: Más de 2.8 millones de aplicaciones auditadas.",
        "Fuentes líderes de telemetría: Veracode, Sonar, Contrast Security, etc.",
        "Análisis inicial de 589 CWEs agrupadas en 248 debilidades finales.",
        "8 categorías surgieron de datos empíricos y 2 por encuesta comunitaria.",
        "Cambio de paradigma: Priorización de la CAUSA RAÍZ por sobre los síntomas."
    ], "Datos Duros", COLOR_BLUE)

    add_card(s4, 6.8, 1.8, 5.7, 5.0, "Grandes Novedades frente a 2021", [
        "A01 (Broken Access Control) se mantiene #1 y absorbe a SSRF.",
        "A02 (Security Misconfiguration) asciende drásticamente del puesto 5 al 2.",
        "NUEVA: A03 Software Supply Chain Failures (refleja el riesgo en compilación y CI/CD).",
        "A04 (Criptografía) y A05 (Inyección) bajan puestos relativos ante fallas arquitectónicas.",
        "NUEVA: A10 Mishandling of Exceptional Conditions (mal manejo de errores y fallas abiertas)."
    ], "Evolución de Amenazas", COLOR_AMBER)

    # =========================================================================
    # SLIDE 5: ORADOR 2 - A01 BROKEN ACCESS CONTROL
    # =========================================================================
    s5 = prs.slides.add_slide(blank_layout)
    apply_bg(s5)
    add_header(s5, "A01: Control de Acceso Defectuoso (Broken Access Control)", "Puesto #1 en OWASP 2025", "Orador 2", COLOR_CRIMSON)

    add_card(s5, 0.8, 1.8, 3.7, 5.0, "Contexto y Alcance", [
        "Mantiene el 1er puesto de forma absoluta.",
        "El 100% de las aplicaciones analizadas tuvo alguna vulnerabilidad de control de acceso.",
        "Ocurre cuando un usuario puede actuar fuera de los límites de sus permisos asignados.",
        "En 2025 absorbe SSRF (Server-Side Request Forgery)."
    ], "100% de Incidencia", COLOR_CRIMSON)

    add_card(s5, 4.8, 1.8, 3.7, 5.0, "Causas Frecuentes & IDOR", [
        "Insecure Direct Object References (IDOR): cambiar un ID en la URL (?acct=102).",
        "APIs sin validación en verbos POST, PUT y DELETE.",
        "Violación del Principio de Mínimo Privilegio.",
        "Tokens JWT o cookies manipuladas para elevar privilegios.",
        "Controles delegados únicamente al navegador/frontend."
    ], "Vectores Típicos", COLOR_AMBER)

    add_card(s5, 4.8 + 4.0, 1.8, 3.7, 5.0, "Estrategia de Mitigación", [
        "Validación estricta de propiedad de recursos en el backend (RBAC / ABAC).",
        "Deshabilitar navegación de directorios y forzar denegación por defecto.",
        "Pruebas unitarias e integración de control de acceso en pipelines CI/CD.",
        "Registro de auditoría ante cualquier intento de elevación de privilegios."
    ], "Defensa en Profundidad", COLOR_EMERALD)

    # =========================================================================
    # SLIDE 6: ORADOR 2 - A02 SECURITY MISCONFIGURATION
    # =========================================================================
    s6 = prs.slides.add_slide(blank_layout)
    apply_bg(s6)
    add_header(s6, "A02: Configuración de Seguridad Incorrecta", "Ascenso al Puesto #2", "Orador 2", COLOR_CRIMSON)

    add_card(s6, 0.8, 1.8, 3.7, 5.0, "¿Por qué subió al Puesto #2?", [
        "Estaba en el puesto #5 en 2021; su escalada refleja la complejidad cloud.",
        "100% de aplicaciones evaluadas tenían configuraciones inseguras.",
        "El comportamiento real del sistema depende más de su configuración que de su código fuente.",
        "Los entornos de nube multiplican la posibilidad de configuraciones erróneas."
    ], "Falla Estructural", COLOR_CRIMSON)

    add_card(s6, 4.8, 1.8, 3.7, 5.0, "Causas Frecuentes", [
        "Cuentas y contraseñas predeterminadas sin modificar.",
        "Puertos, consolas de administración o páginas de prueba habilitadas.",
        "Headers de seguridad HTTP ausentes (CSP, HSTS, X-Content-Type).",
        "Permisos excesivos en buckets y servicios cloud.",
        "Mensajes de error verbosos que exponen stack traces y versiones."
    ], "Errores Comunes", COLOR_AMBER)

    add_card(s6, 8.8, 1.8, 3.7, 5.0, "Mitigación y Detección", [
        "Guías de endurecimiento (Hardening) basadas en estándares NIST y CIS.",
        "Infraestructura como Código (IaC) auditada de forma reproducible.",
        "Deshabilitar características y servicios innecesarios en imágenes Docker.",
        "Verificaciones automatizadas comparando contra una línea base segura."
    ], "Línea Base Segura", COLOR_EMERALD)

    # =========================================================================
    # SLIDE 7: ORADOR 3 - A03 SOFTWARE SUPPLY CHAIN FAILURES
    # =========================================================================
    s7 = prs.slides.add_slide(blank_layout)
    apply_bg(s7)
    add_header(s7, "A03: Fallas en la Cadena de Suministro de Software", "Nueva Categoría 2025 (#3)", "Orador 3", COLOR_AMBER)

    add_card(s7, 0.8, 1.8, 3.7, 5.0, "Evolución de la Amenaza", [
        "Reemplaza y amplía 'componentes vulnerables y desactualizados'.",
        "Abarca todo el ciclo: herramientas de build, repositorios, dependencias y CI/CD.",
        "Top #1 en la votación comunitaria de expertos por su impacto devastador.",
        "Los atacantes ya no atacan tu app: atacan las dependencias de tus dependencias."
    ], "Nueva Categoría", COLOR_AMBER)

    add_card(s7, 4.8, 1.8, 3.7, 5.0, "Caso Real: SolarWinds (2019)", [
        "Atacantes comprometieron el pipeline de compilación de la empresa.",
        "El malware se distribuyó firmado y certificado a clientes legítimos.",
        "Afectó a más de 18.000 organizaciones globales, incluyendo el gobierno de EEUU y Microsoft.",
        "Demostró que una sola dependencia comprometida vulnera todo un ecosistema."
    ], "Impacto Sistémico", COLOR_CRIMSON)

    add_card(s7, 8.8, 1.8, 3.7, 5.0, "Controles y Mitigación", [
        "Generación y monitoreo continuo de SBOM (Software Bill of Materials).",
        "Herramientas activas: OWASP Dependency-Track y Dependency-Check.",
        "Consultas periódicas a bases de vulnerabilidades (NVD y OSV).",
        "Segregación estricta de funciones y protección del pipeline de CI/CD."
    ], "Visibilidad de Código", COLOR_EMERALD)

    # =========================================================================
    # SLIDE 8: ORADOR 3 - A04 CRYPTOGRAPHIC FAILURES
    # =========================================================================
    s8 = prs.slides.add_slide(blank_layout)
    apply_bg(s8)
    add_header(s8, "A04: Fallas Criptográficas", "Puesto #4 en OWASP 2025", "Orador 3", COLOR_AMBER)

    add_card(s8, 0.8, 1.8, 3.7, 5.0, "Alcance del Problema", [
        "Anteriormente conocida como Exposición de Datos Sensibles.",
        "Agrupa problemas de ausencia de cifrado, cifrado débil o filtración de claves.",
        "Punto crítico: generadores de números pseudoaleatorios inadecuados.",
        "Resultado habitual: compromiso total de confidencialidad de datos."
    ], "Protección de Datos", COLOR_AMBER)

    add_card(s8, 4.8, 1.8, 3.7, 5.0, "Causas Frecuentes", [
        "Uso de algoritmos obsoletos o rotos (ej: MD5, SHA1).",
        "Tráfico en texto plano sin forzar HTTPS/TLS en toda la aplicación.",
        "Almacenamiento de contraseñas con algoritmos simples o sin salt.",
        "Claves de cifrado harcodeadas en repositorios de Git públicos o privados.",
        "Falta de políticas de rotación de claves criptográficas."
    ], "Errores Clave", COLOR_CRIMSON)

    add_card(s8, 8.8, 1.8, 3.7, 5.0, "Recomendaciones Técnicas", [
        "TLS configurado y forzado en todas las rutas con certificados válidos.",
        "Clasificación formal de datos según su sensibilidad en reposo y en tránsito.",
        "Funciones de hash diseñadas para contraseñas: Argon2, bcrypt o PBKDF2.",
        "Gestores de secretos corporativos (HashiCorp Vault, AWS Secrets Manager)."
    ], "Buenas Prácticas", COLOR_EMERALD)

    # =========================================================================
    # SLIDE 9: ORADOR 4 - A05 INJECTION
    # =========================================================================
    s9 = prs.slides.add_slide(blank_layout)
    apply_bg(s9)
    add_header(s9, "A05: Inyección (SQL Injection & XSS)", "Puesto #5 — Récord de CVEs", "Orador 4", COLOR_PURPLE)

    add_card(s9, 0.8, 1.8, 3.7, 5.0, "Magnitud de la Inyección", [
        "Mayor volumen histórico de CVEs: más de 44.000 vulnerabilidades reportadas.",
        "XSS (Cross-Site Scripting): Alta frecuencia, más de 30.000 casos.",
        "SQLi (Inyección SQL): Menor frecuencia pero con impacto devastador (14.000 casos).",
        "Ocurre cuando datos no confiables se ejecutan como instrucciones de código."
    ], "+44k CVEs", COLOR_PURPLE)

    add_card(s9, 4.8, 1.8, 3.7, 5.0, "Causa Raíz: Concatenación", [
        "Construcción de consultas SQL mediante concatenación directa de strings.",
        "Ejemplo: 'SELECT * FROM users WHERE user = ' + input + ...",
        "Payloads booleanos como ' OR '1'='1 alteran el AST del intérprete SQL.",
        "Se anulan las condiciones de autenticación, otorgando acceso como administrador."
    ], "Mecanismo del Ataque", COLOR_CRIMSON)

    add_card(s9, 8.8, 1.8, 3.7, 5.0, "Solución: Prepared Statements", [
        "Consultas parametrizadas obligatorias (los parámetros viajan como datos, no código).",
        "Validación y saneamiento de entradas con esquemas estrictos (Zod, Joi).",
        "Integración de análisis estático y dinámico (SAST, DAST, IAST) en CI/CD.",
        "Evitar consultas 'raw' dinámicas incluso al utilizar frameworks u ORMs."
    ], "Mitigación Definitiva", COLOR_EMERALD)

    # =========================================================================
    # SLIDE 10: ORADOR 4 - A06 INSECURE DESIGN
    # =========================================================================
    s10 = prs.slides.add_slide(blank_layout)
    apply_bg(s10)
    add_header(s10, "A06: Diseño Inseguro (Insecure Design)", "Puesto #6 en OWASP 2025", "Orador 4", COLOR_PURPLE)

    add_card(s10, 0.8, 1.8, 3.7, 5.0, "Diseño vs. Implementación", [
        "Principio clave de OWASP: Separar el diseño de la implementación.",
        "Un diseño inseguro no se soluciona con una implementación perfecta.",
        "Si los controles no fueron concebidos desde la arquitectura, el código fallará.",
        "Diferencia: Un bug es un error de código; una falla de diseño es una ausencia de lógica."
    ], "Conceptos > Código", COLOR_PURPLE)

    add_card(s10, 4.8, 1.8, 3.7, 5.0, "Ejemplo Real: Lógica de Negocio", [
        "Caso del sistema de cine: descuentos grupales a partir de 15 personas.",
        "Atacante reserva 600 asientos en todas las salas con pocas peticiones.",
        "Causa un daño financiero inmenso sin romper ninguna validación técnica de código.",
        "Falla de diseño: no contemplar límites de tasa o retención temporal por usuario."
    ], "Falla de Lógica", COLOR_CRIMSON)

    add_card(s10, 8.8, 1.8, 3.7, 5.0, "Enfoque Preventivo", [
        "Modelado de Amenazas (Threat Modeling) en las fases tempranas del ciclo.",
        "Requerimientos de seguridad explícitos en las historias de usuario.",
        "Diseño de casos de abuso (Abuse Cases) junto a los casos de uso habituales.",
        "Implementación del marco de madurez OWASP SAMM."
    ], "Arquitectura Segura", COLOR_EMERALD)

    # =========================================================================
    # SLIDE 11: ORADOR 5 - A07 & A08
    # =========================================================================
    s11 = prs.slides.add_slide(blank_layout)
    apply_bg(s11)
    add_header(s11, "A07: Autenticación & A08: Integridad de Software", "Puestos #7 y #8", "Orador 5", COLOR_EMERALD)

    add_card(s11, 0.8, 1.8, 5.7, 5.0, "A07: Fallas de Autenticación", [
        "Ocurre cuando el sistema reconoce erróneamente a un usuario ilegítimo.",
        "Ataques típicos: Credential Stuffing y Fuerza Bruta automatizada.",
        "El problema humano: de 'Winter2025' a 'Winter2026' tras filtraciones.",
        "Sin rate limiting, el backend opera como un 'oráculo' para el botnet.",
        "Mitigación: Pautas NIST 800-63b, bloqueo por intentos y MFA obligatorio."
    ], "Puesto #7", COLOR_AMBER)

    add_card(s11, 6.8, 1.8, 5.7, 5.0, "A08: Fallas de Integridad de Software y Datos", [
        "Asumir que código o artefactos externos son confiables sin verificación.",
        "Pipelines que descargan módulos sin verificar firma criptográfica.",
        "Actualizaciones automáticas aplicadas a ciegas.",
        "Deserialización insegura: manipulación de objetos para ejecutar código remoto.",
        "Mitigación: Firmas digitales, repositorios espejo auditados y escáneres estáticos."
    ], "Puesto #8", COLOR_EMERALD)

    # =========================================================================
    # SLIDE 12: ORADOR 5 - A09 & A10
    # =========================================================================
    s12 = prs.slides.add_slide(blank_layout)
    apply_bg(s12)
    add_header(s12, "A09: Registro y Alertas & A10: Condiciones Excepcionales", "Puestos #9 y #10 (Novedad)", "Orador 5", COLOR_EMERALD)

    add_card(s12, 0.8, 1.8, 5.7, 5.0, "A09: Fallas de Registro y Alertas", [
        "Evolución clave: Pasó de llamarse 'monitoreo' a 'alertas'.",
        "Tener logs sin alertas no sirve: ataques persisten años inadvertidos.",
        "Caso Real: Filtración de aseguradora de salud infantil durante 7 años (3.5M de niños).",
        "Mitigación: Pila ELK, honeytokens (trampas con cero falsos positivos) y plan NIST 800-61."
    ], "Puesto #9", COLOR_BLUE)

    add_card(s12, 6.8, 1.8, 5.7, 5.0, "A10: Mal Manejo de Condiciones Excepcionales", [
        "NUEVA CATEGORÍA EN 2025: 24 debilidades enfocadas en errores.",
        "Falla abierta: el sistema falla dejando datos o permisos expuestos.",
        "Ejemplo: Transacción bancaria en 3 pasos que no ejecuta rollback ante cortes.",
        "Mitigación: Manejo centralizado global de excepciones, idempotencia y límites de cuota."
    ], "Puesto #10 (Nueva)", COLOR_CRIMSON)

    # =========================================================================
    # SLIDE 13: CONCLUSIÓN & DEMO
    # =========================================================================
    s13 = prs.slides.add_slide(blank_layout)
    apply_bg(s13)
    add_header(s13, "Conclusiones del Estudio & Demostración Práctica", "Cierre y Laboratorio en Vivo", "Equipo Completo", COLOR_CYAN)

    add_card(s13, 0.8, 1.8, 5.7, 5.0, "Conclusiones de la Monografía", [
        "1. La seguridad no se añade al final: debe ser parte del diseño inicial.",
        "2. El Top 10:2025 destaca que el código bien escrito puede ser vulnerado por infraestructura y cadena de suministro.",
        "3. La automatización ofensiva exige automatización defensiva (SAST, DAST, SBOM y MFA).",
        "4. Enfoque integral: Diseño Seguro + Automatización CI/CD + Monitoreo con Alertas."
    ], "Pilares Clave", COLOR_CYAN)

    add_card(s13, 6.8, 1.8, 5.7, 5.0, "Demostración del Simulador Web (SeguridadTP)", [
        "Laboratorio interactivo desarrollado para acompañar la monografía:",
        "• Demo 1: Explotación y mitigación de IDOR (A01 - Broken Access Control).",
        "• Demo 2: Bypass de autenticación con SQL Injection y Prepared Statements (A05).",
        "• Demo 3: Simulación de botnet de fuerza bruta, bloqueo NIST y MFA (A07).",
        "Disponible en: https://github.com/Vgarzaniti/SeguridadTP"
    ], "Simulación en Vivo", COLOR_EMERALD)

    output_path = "OWASP_Top_10_2025_Presentacion.pptx"
    prs.save(output_path)
    print(f"Presentation generated successfully: {output_path}")

if __name__ == "__main__":
    create_deck()
